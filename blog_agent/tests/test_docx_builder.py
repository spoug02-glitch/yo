from docx import Document
from docx.shared import Pt

from blog_agent import docx_builder


def test_build_docx_has_titles_headings_and_font(tmp_path):
    body = "# AICE 가이드\n\n## 1. 개념\n\n본문 문장입니다.\n\n{{chart:t1}}\n"
    table_specs = {"t1": {"data": {"headers": ["등급", "응시료"],
                                    "rows": [["Basic", "50,000원"], ["Associate", "80,000원"]]}}}
    out = docx_builder.build_docx(
        title_candidates=["제목 후보 A", "제목 후보 B"],
        body_md=body, table_specs=table_specs, chart_images={},
        sources=["공식 홈페이지 aice.study — 확인일 2026-07-29"],
        out_path=tmp_path / "out.docx",
    )
    assert out.exists()
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert any("제목 후보 A" in t for t in texts)
    assert any("AICE 가이드" == t for t in texts)
    # 소제목 19pt 확인
    h2 = next(p for p in doc.paragraphs if p.text == "1. 개념")
    assert h2.runs[0].font.size == Pt(19)
    assert h2.runs[0].font.bold is True
    # 네이티브 표가 삽입됐는지
    assert len(doc.tables) == 1
    assert doc.tables[0].cell(0, 0).text == "등급"


def test_body_font_is_nanum(tmp_path):
    out = docx_builder.build_docx(
        title_candidates=["t"], body_md="본문", table_specs={}, chart_images={},
        sources=[], out_path=tmp_path / "o.docx")
    doc = Document(str(out))
    body_p = next(p for p in doc.paragraphs if p.text == "본문")
    assert body_p.runs[0].font.name == "나눔고딕"


def test_inline_bold_becomes_bold_run(tmp_path):
    out = docx_builder.build_docx(
        title_candidates=["t"], body_md="일반 **강조** 꼬리", table_specs={}, chart_images={},
        sources=[], out_path=tmp_path / "b.docx")
    doc = Document(str(out))
    p = next(p for p in doc.paragraphs if p.text == "일반 강조 꼬리")
    runs = [(r.text, r.font.bold) for r in p.runs]
    assert ("강조", True) in runs
    assert all(b is not True for t, b in runs if t != "강조")


def test_unresolved_chart_marker_shows_visible_placeholder(tmp_path):
    body = "본문\n\n{{chart:missing_id}}\n"
    out = docx_builder.build_docx(
        title_candidates=["t"], body_md=body, table_specs={}, chart_images={},
        sources=[], out_path=tmp_path / "p.docx")
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert "[도표 누락: missing_id]" in texts
