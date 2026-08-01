from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches

_FONT = "나눔고딕"
_CHART_RE = re.compile(r"^\{\{chart:(?P<id>[^}]+)\}\}$")
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _style_run(run, size_pt: int, *, bold=False, italic=False):
    run.font.name = _FONT
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    # 한글 폰트 East Asian 힌트
    run._element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", _FONT)


def _add_para(doc, text: str, size_pt: int, *, bold=False, italic=False):
    p = doc.add_paragraph()
    _style_run(p.add_run(text), size_pt, bold=bold, italic=italic)
    return p


def _add_rich_para(doc, text: str, size_pt: int, *, bold=False, italic=False, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    pos = 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            _style_run(p.add_run(text[pos:m.start()]), size_pt, bold=bold, italic=italic)
        _style_run(p.add_run(m.group(1)), size_pt, bold=True, italic=italic)
        pos = m.end()
    if pos < len(text):
        _style_run(p.add_run(text[pos:]), size_pt, bold=bold, italic=italic)
    return p


def _add_table(doc, data: dict):
    headers, rows = data["headers"], data["rows"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        _style_run(table.cell(0, i).paragraphs[0].add_run(str(h)), 15, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            _style_run(cells[i].paragraphs[0].add_run(str(val)), 15)


def build_docx(*, title_candidates: list[str], body_md: str,
               table_specs: dict[str, dict], chart_images: dict[str, Path],
               sources: list[str], out_path: Path) -> Path:
    doc = Document()

    _add_para(doc, "제목 후보", 19, bold=True)
    for i, t in enumerate(title_candidates, 1):
        _add_para(doc, f"{i}. {t}", 15)
    doc.add_paragraph()

    for line in body_md.splitlines():
        m = _CHART_RE.match(line.strip())
        if m:
            cid = m.group("id")
            if cid in table_specs:
                _add_table(doc, table_specs[cid]["data"])
            elif cid in chart_images:
                doc.add_picture(str(chart_images[cid]), width=Inches(5.5))
            else:
                _add_para(doc, f"[도표 누락: {cid}]", 15)
            continue
        if line.startswith("### "):
            _add_para(doc, line[4:], 15, italic=True)
        elif line.startswith("## "):
            _add_para(doc, line[3:], 19, bold=True)
        elif line.startswith("# "):
            _add_para(doc, line[2:], 24, bold=True)
        elif line.startswith("- "):
            _add_rich_para(doc, line[2:], 15, style="List Bullet")
        elif line.strip():
            _add_rich_para(doc, line, 15)

    if sources:
        _add_para(doc, "출처", 19, bold=True)
        for s in sources:
            _add_para(doc, s, 15)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
